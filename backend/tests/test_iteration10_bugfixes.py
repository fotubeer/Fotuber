"""
Fotuber Iteration 10 backend tests: multi-bug-fix session verification

Coverage:
  1. PDF export (/api/transactions/export.pdf) — Turkish chars and ₺ symbol
     render correctly (using DejaVu Sans font). Extract PDF text with pypdf
     and verify presence of 'Nakit Akışı Raporu', 'Dönem', 'Tür', 'Açıklama',
     '₺', and a Turkish transaction description.
  2. Excel export (/api/transactions/export.xlsx) — regression: Turkish headers
     and Turkish content still correct.
  3. POST /api/admin/photo-albums/{aid}/photos — multipart upload of 3 JPGs
     without manual Content-Type header, photo_count increments.
  4. POST /api/guest-events/{token}/upload — customer uploads a small file
     without manual Content-Type header, kvkk_accepted=true.
"""
import io
import os
import time
import uuid
from datetime import datetime, timezone

import pytest
import requests
from PIL import Image
from openpyxl import load_workbook
from pypdf import PdfReader


# ---------------------------------------------------------------------------
# Backend URL
# ---------------------------------------------------------------------------
def _load_backend_url() -> str:
    url = os.environ.get("REACT_APP_BACKEND_URL")
    if not url:
        env_path = "/app/frontend/.env"
        if os.path.exists(env_path):
            with open(env_path) as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        url = line.split("=", 1)[1].strip().strip('"').strip("'")
                        break
    assert url, "REACT_APP_BACKEND_URL must be set"
    return url.rstrip("/")


BASE_URL = _load_backend_url()
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@fotuber.com.tr"
ADMIN_PASSWORD = "FTB.2024"
TODAY = datetime.now(timezone.utc).date().isoformat()

TURKISH_DESC = "Düğün çekimi — İstanbul dış çekim ğüşöçİĞÜŞÖÇ"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _make_jpeg(size=(400, 300), color=(120, 200, 90)) -> bytes:
    im = Image.new("RGB", size, color)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=80)
    return b.getvalue()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def admin_client():
    s = requests.Session()
    r = s.post(f"{API}/auth/login",
               json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
               timeout=30)
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    token = r.json()["token"]
    s.headers.update({"Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="module")
def customer_client():
    email = f"test_it10_{uuid.uuid4().hex[:8]}@example.com"
    s = requests.Session()
    r = s.post(f"{API}/auth/register", json={
        "name": "Guest",
        "surname": "Uploader",
        "email": email,
        "phone": "05001112233",
        "password": "TestPass1!",
        "kvkk_consent": True,
        "marketing_consent": False,
    }, timeout=30)
    assert r.status_code == 200, f"register failed: {r.status_code} {r.text}"
    token = r.json()["token"]
    s.headers.update({"Authorization": f"Bearer {token}"})
    s.test_email = email
    return s


# ---------------------------------------------------------------------------
# Bug 4: PDF export with Turkish + ₺
# ---------------------------------------------------------------------------
class TestPdfTurkishFont:
    """Verify DejaVu font registration → Turkish chars + ₺ render (not empty boxes)."""

    @pytest.fixture(scope="class")
    def seeded_turkish_tx(self, admin_client):
        # Insert a transaction with Turkish description for TODAY
        r = admin_client.post(f"{API}/transactions", json={
            "kind": "income",
            "amount": 4567.89,
            "payment_method": "cash",
            "category": "TEST_Kategori_ç",
            "description": TURKISH_DESC,
            "date": TODAY,
        }, timeout=30)
        assert r.status_code == 200, f"seed tx failed: {r.status_code} {r.text}"
        tid = r.json()["id"]
        yield tid
        admin_client.delete(f"{API}/transactions/{tid}", timeout=30)

    def test_pdf_contains_turkish_and_lira(self, admin_client, seeded_turkish_tx):
        r = admin_client.get(f"{API}/transactions/export.pdf",
                             params={"period": "today"}, timeout=60)
        assert r.status_code == 200, f"pdf export failed: {r.status_code} {r.text[:200]}"
        assert r.content[:5] == b"%PDF-"

        # Extract text
        reader = PdfReader(io.BytesIO(r.content))
        text = "\n".join([(p.extract_text() or "") for p in reader.pages])
        assert text, "PDF has no extractable text"

        # Required substrings (Turkish + ₺)
        must_have = [
            "Nakit Akışı Raporu",   # Title with 'ı'
            "Dönem",                # 'ö'
            "Tür",                  # 'ü'
            "Açıklama",             # 'ç'
            "Kategori",
            "₺",                    # Turkish Lira symbol
            "Düğün",                # 'ü', 'ğ'
            "İstanbul",             # dotted uppercase İ
        ]
        missing = [s for s in must_have if s not in text]
        assert not missing, (
            f"PDF missing Turkish substrings: {missing}. "
            f"First 500 chars of extracted text: {text[:500]!r}"
        )

    def test_pdf_valid_pdf_all_periods(self, admin_client, seeded_turkish_tx):
        for period in ["today", "month", "year", "all"]:
            r = admin_client.get(f"{API}/transactions/export.pdf",
                                 params={"period": period}, timeout=60)
            assert r.status_code == 200, f"[{period}] {r.status_code}"
            assert r.content[:5] == b"%PDF-", f"[{period}] not a pdf"


# ---------------------------------------------------------------------------
# Regression: XLSX still ok
# ---------------------------------------------------------------------------
class TestXlsxRegression:
    def test_xlsx_turkish_headers(self, admin_client):
        r = admin_client.get(f"{API}/transactions/export.xlsx",
                             params={"period": "month"}, timeout=60)
        assert r.status_code == 200
        assert r.content[:4] == b"PK\x03\x04"
        wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        ws = wb.active
        row4 = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
        row4 = [v for v in row4 if v is not None]
        assert "Tarih" in row4
        assert "Tür" in row4
        assert "Açıklama" in row4
        # Tutar (₺) - the ₺ char should still be in xlsx too
        assert any(v and "₺" in str(v) for v in row4), f"₺ not in xlsx headers: {row4}"


# ---------------------------------------------------------------------------
# Bug 1/2: photo album upload — no manual Content-Type
# ---------------------------------------------------------------------------
class TestAlbumUpload:
    @pytest.fixture(scope="class")
    def test_album(self, admin_client):
        r = admin_client.post(f"{API}/admin/photo-albums", json={
            "couple_names": "TEST_Iter10_Album",
            "event_date": TODAY,
            "notes": "",
            "max_selections": 50,
        }, timeout=30)
        assert r.status_code == 200, f"create album failed: {r.status_code} {r.text}"
        album = r.json()
        yield album
        admin_client.delete(f"{API}/admin/photo-albums/{album['id']}", timeout=30)

    def test_upload_three_photos(self, admin_client, test_album):
        aid = test_album["id"]
        jpg_bytes = _make_jpeg((640, 480), (30, 120, 210))
        files = [
            ("files", (f"TEST_photo_{i}.jpg", jpg_bytes, "image/jpeg"))
            for i in range(3)
        ]
        # multipart via requests → no manual Content-Type set (mimics the frontend fix)
        r = admin_client.post(
            f"{API}/admin/photo-albums/{aid}/photos",
            files=files,
            timeout=60,
        )
        assert r.status_code == 200, f"upload failed: {r.status_code} {r.text[:300]}"
        data = r.json()
        assert data.get("added") == 3, f"expected added=3, got {data}"
        assert len(data.get("photos", [])) == 3

        # Verify photo_count in album
        g = admin_client.get(f"{API}/admin/photo-albums/{aid}", timeout=30)
        assert g.status_code == 200
        assert g.json()["album"]["photo_count"] == 3


# ---------------------------------------------------------------------------
# Bug 1/2: guest event upload — no manual Content-Type
# ---------------------------------------------------------------------------
class TestGuestEventUpload:
    @pytest.fixture(scope="class")
    def test_event(self, admin_client):
        r = admin_client.post(f"{API}/admin/guest-events", json={
            "name": "TEST_Iter10_Event",
            "couple_names": "TEST_Ada & Bora",
            "event_date": TODAY,
            "welcome_message": "Hoş geldiniz",
            "max_size_per_user_mb": 50,
            "retention_days": 3,
        }, timeout=30)
        assert r.status_code == 200, f"create event failed: {r.status_code} {r.text}"
        ev = r.json()
        yield ev
        admin_client.delete(f"{API}/admin/guest-events/{ev['id']}", timeout=30)

    def test_customer_upload_file(self, customer_client, test_event):
        token = test_event["upload_token"]
        jpg_bytes = _make_jpeg((300, 200), (255, 200, 100))
        files = {"file": ("TEST_guest.jpg", jpg_bytes, "image/jpeg")}
        data = {"kvkk_accepted": "true"}
        r = customer_client.post(
            f"{API}/guest-events/{token}/upload",
            files=files,
            data=data,
            timeout=60,
        )
        assert r.status_code == 200, f"guest upload failed: {r.status_code} {r.text[:300]}"
        body = r.json()
        assert body.get("ok") is True
        assert body.get("size") == len(jpg_bytes)
        assert "remaining" in body
        assert body["remaining"] >= 0
