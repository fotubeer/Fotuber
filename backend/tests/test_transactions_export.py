"""
Fotuber Iteration 6 backend tests: Cash-flow Transactions Export (XLSX / PDF)

Coverage:
  - Excel (.xlsx) export for periods: today, week, month, year, all
  - PDF (.pdf) export for periods: today, week, month, year, all
  - Auth guards: unauthenticated -> 401, customer -> 403
  - Response headers: Content-Type, Content-Disposition (attachment, filename)
  - File signature: XLSX zip magic PK\x03\x04, PDF magic %PDF-
  - openpyxl parse to verify header row (row 4) matches Turkish column labels
  - Data rows below reflect inserted transaction
"""
import io
import os
import uuid
from datetime import datetime, timezone

import pytest
import requests
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Backend URL from frontend/.env (or REACT_APP_BACKEND_URL)
# ---------------------------------------------------------------------------
def _load_backend_url() -> str:
    url = os.environ.get("REACT_APP_BACKEND_URL")
    if not url:
        env_path = "/app/frontend/.env"
        if os.path.exists(env_path):
            with open(env_path) as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        url = line.split("=", 1)[1].strip().strip('"').strip("'")
                        break
    assert url, "REACT_APP_BACKEND_URL must be set"
    return url.rstrip("/")


BASE_URL = _load_backend_url()
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@fotuber.com.tr"
ADMIN_PASSWORD = "FTB.2024"

TODAY = datetime.now(timezone.utc).date().isoformat()

EXPECTED_HEADERS = [
    "Tarih",
    "Tür",
    "Ödeme Yöntemi",
    "Kategori",
    "Açıklama",
    "Tutar (₺)",
    "İşaretli Tutar",
]

PERIODS = ["today", "week", "month", "year", "all"]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def admin_client():
    s = requests.Session()
    r = s.post(
        f"{API}/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        timeout=30,
    )
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    token = r.json()["token"]
    s.headers.update({"Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="module")
def customer_client():
    email = f"TEST_export_{uuid.uuid4().hex[:8]}@example.com"
    s = requests.Session()
    r = s.post(
        f"{API}/auth/register",
        json={
            "name": "Export",
            "surname": "Guard",
            "email": email,
            "phone": "05001112233",
            "password": "pass1234",
            "kvkk_consent": True,
            "marketing_consent": False,
        },
        timeout=30,
    )
    assert r.status_code == 200, f"register failed: {r.status_code} {r.text}"
    token = r.json()["token"]
    s.headers.update({"Authorization": f"Bearer {token}"})
    return s


@pytest.fixture(scope="module")
def seeded_transaction(admin_client):
    """Insert one INCOME + one EXPENSE for today. Cleanup after module."""
    created_ids = []
    # Clean today's existing tx first (to keep assertions predictable)
    r = admin_client.get(
        f"{API}/transactions",
        params={"date_from": TODAY, "date_to": TODAY},
        timeout=30,
    )
    if r.status_code == 200:
        for t in r.json():
            admin_client.delete(f"{API}/transactions/{t['id']}", timeout=30)

    payloads = [
        {
            "kind": "income",
            "amount": 1234.56,
            "payment_method": "cash",
            "category": "TEST_Kategori",
            "description": "TEST_export_income",
            "date": TODAY,
        },
        {
            "kind": "expense",
            "amount": 78.90,
            "payment_method": "card",
            "category": "TEST_Gider",
            "description": "TEST_export_expense",
            "date": TODAY,
        },
    ]
    for p in payloads:
        r = admin_client.post(f"{API}/transactions", json=p, timeout=30)
        assert r.status_code == 200, f"create tx failed: {r.status_code} {r.text}"
        created_ids.append(r.json()["id"])

    yield payloads

    for tid in created_ids:
        admin_client.delete(f"{API}/transactions/{tid}", timeout=30)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _assert_attachment_header(headers, ext):
    cd = headers.get("Content-Disposition") or headers.get("content-disposition") or ""
    assert "attachment" in cd.lower(), f"Content-Disposition missing 'attachment': {cd!r}"
    assert ext in cd.lower(), f"Content-Disposition missing '{ext}': {cd!r}"


# ---------------------------------------------------------------------------
# Excel export tests
# ---------------------------------------------------------------------------
class TestExcelExport:
    @pytest.mark.parametrize("period", PERIODS)
    def test_xlsx_all_periods(self, admin_client, seeded_transaction, period):
        r = admin_client.get(
            f"{API}/transactions/export.xlsx",
            params={"period": period},
            timeout=60,
        )
        assert r.status_code == 200, f"[{period}] {r.status_code} {r.text[:200]}"
        assert r.headers.get("content-type", "").startswith(XLSX_MIME), (
            f"[{period}] bad content-type: {r.headers.get('content-type')}"
        )
        # ZIP magic PK\x03\x04
        assert r.content[:4] == b"PK\x03\x04", (
            f"[{period}] not a valid xlsx (magic={r.content[:4]!r})"
        )
        _assert_attachment_header(r.headers, ".xlsx")

    def test_xlsx_month_header_row_and_data(self, admin_client, seeded_transaction):
        r = admin_client.get(
            f"{API}/transactions/export.xlsx",
            params={"period": "month"},
            timeout=60,
        )
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        ws = wb.active
        # Row 4 = headers
        row4 = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
        # trim trailing None cols
        row4 = row4[: len(EXPECTED_HEADERS)]
        assert row4 == EXPECTED_HEADERS, f"unexpected headers: {row4}"

        # Data rows (row 5+): scan first ~50 rows and verify seeded tx present
        data_rows = []
        for row in ws.iter_rows(min_row=5, max_row=60, values_only=True):
            if not row or all(v in (None, "") for v in row):
                continue
            data_rows.append(row)

        # Look for our seeded income row (Tarih=TODAY, Tür=Gelir, Yöntem=Nakit, amount=1234.56)
        matches = [
            r_ for r_ in data_rows
            if r_[0] == TODAY
            and r_[1] == "Gelir"
            and r_[2] == "Nakit"
            and r_[4] == "TEST_export_income"
            and float(r_[5]) == 1234.56
            and float(r_[6]) == 1234.56  # signed = +amount for income
        ]
        assert matches, (
            "Seeded income row not found in xlsx data rows. "
            f"Rows sample: {data_rows[:5]}"
        )

        # And the expense row (signed = -78.90)
        exp_matches = [
            r_ for r_ in data_rows
            if r_[0] == TODAY
            and r_[1] == "Gider"
            and r_[2] == "Kart"
            and r_[4] == "TEST_export_expense"
            and float(r_[5]) == 78.90
            and float(r_[6]) == -78.90
        ]
        assert exp_matches, (
            "Seeded expense row not found in xlsx data rows. "
            f"Rows sample: {data_rows[:5]}"
        )


# ---------------------------------------------------------------------------
# PDF export tests
# ---------------------------------------------------------------------------
class TestPdfExport:
    @pytest.mark.parametrize("period", PERIODS)
    def test_pdf_all_periods(self, admin_client, seeded_transaction, period):
        r = admin_client.get(
            f"{API}/transactions/export.pdf",
            params={"period": period},
            timeout=60,
        )
        assert r.status_code == 200, f"[{period}] {r.status_code} {r.text[:200]}"
        assert r.headers.get("content-type", "").startswith(PDF_MIME), (
            f"[{period}] bad content-type: {r.headers.get('content-type')}"
        )
        assert r.content[:5] == b"%PDF-", (
            f"[{period}] not a valid pdf (magic={r.content[:8]!r})"
        )
        _assert_attachment_header(r.headers, ".pdf")


# ---------------------------------------------------------------------------
# Auth guard tests
# ---------------------------------------------------------------------------
class TestExportAuth:
    def test_xlsx_no_auth_returns_401(self):
        r = requests.get(
            f"{API}/transactions/export.xlsx",
            params={"period": "month"},
            timeout=30,
        )
        assert r.status_code == 401, f"expected 401, got {r.status_code}: {r.text[:200]}"

    def test_pdf_no_auth_returns_401(self):
        r = requests.get(
            f"{API}/transactions/export.pdf",
            params={"period": "year"},
            timeout=30,
        )
        assert r.status_code == 401, f"expected 401, got {r.status_code}: {r.text[:200]}"

    def test_xlsx_customer_returns_403(self, customer_client):
        r = customer_client.get(
            f"{API}/transactions/export.xlsx",
            params={"period": "month"},
            timeout=30,
        )
        assert r.status_code == 403, f"expected 403, got {r.status_code}: {r.text[:200]}"

    def test_pdf_customer_returns_403(self, customer_client):
        r = customer_client.get(
            f"{API}/transactions/export.pdf",
            params={"period": "year"},
            timeout=30,
        )
        assert r.status_code == 403, f"expected 403, got {r.status_code}: {r.text[:200]}"
